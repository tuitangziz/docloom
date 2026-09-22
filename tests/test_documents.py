from io import BytesIO

from docx import Document as Word
from pypdf import PdfWriter
import pytest
from reportlab.pdfgen.canvas import Canvas

from docloom.ingest import DocumentError, display_name, parse_document, split_text
from docloom.library import Library


def test_text_location_name_and_chunk_coverage():
    doc = parse_document(r'C:\private\notes.md', b'alpha paragraph\n\nbeta paragraph\n')
    assert doc.filename == 'notes.md'
    assert '起始行 3' in doc.passages[1].location
    text = ''.join(chr(0x4E00 + i) for i in range(2400))
    chunks = split_text(text, 900, 120)
    assert all(len(chunk) <= 900 for chunk in chunks)
    assert all(char in ''.join(chunks) for char in text)
    assert chunks[0][-120:] == chunks[1][:120]
    assert display_name('../../\x00name.txt') == 'name.txt'


def test_pdf_pages_and_docx_tables():
    stream = BytesIO()
    pdf = Canvas(stream)
    pdf.drawString(40, 700, 'First page: orchard opens at nine.')
    pdf.showPage()
    pdf.drawString(40, 700, 'Second page: workshop closes at five.')
    pdf.save()
    doc = parse_document('manual.pdf', stream.getvalue())
    assert 'PDF 第 2 页' in doc.passages[1].location
    assert 'workshop' in doc.passages[1].text
    word = Word()
    word.add_paragraph('预约需提前一天。')
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = '设备', '显卡'
    stream = BytesIO()
    word.save(stream)
    doc = parse_document('manual.docx', stream.getvalue())
    assert '段落 1' in doc.passages[0].location
    assert '表格 1 · 行 1' in doc.passages[1].location
    assert '设备 | 显卡' == doc.passages[1].text


@pytest.mark.parametrize('name,data', [('bad.txt', b'\xff\xfe'), ('bad.txt', b'abc\x00def'), ('bad.docx', b'not a zip'), ('bad.pdf', b'not a pdf'), ('bad.exe', b'hello'), ('empty.md', b'')])
def test_invalid_inputs(name, data):
    with pytest.raises(DocumentError):
        parse_document(name, data)


def test_encrypted_and_scan_only_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(DocumentError, match='OCR'):
        parse_document('scan.pdf', stream.getvalue())
    writer.encrypt('invented-test-password')
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(DocumentError, match='加密'):
        parse_document('locked.pdf', stream.getvalue())


def test_dedup_deletion_and_session_isolation():
    a, b = Library(), Library()
    doc, added = a.add('a.txt', b'orchard peaches available')
    assert added
    assert not a.add('copy.txt', b'orchard peaches available')[1]
    assert len(a.documents) == 1 and not b.documents
    a.remove({doc.id})
    assert a.index is None and not a.documents


def test_tiny_document_can_be_indexed():
    library = Library()
    library.add('tiny.txt', b'x')
    assert library.index.search('x')
